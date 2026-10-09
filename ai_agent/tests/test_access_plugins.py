"""Access rules for skills and plugins (AI Settings), and the plugin export tools."""

import os
from unittest.mock import patch

import frappe
from frappe.tests import IntegrationTestCase

from ai_agent.runtime import access
from ai_agent.tests.utils import delete_docs, make_test_run

USER = "rules-test@example.com"


def rule(name, allow=1, role="", user="", doctype=access.SKILL):
    return {"doctype": doctype, "name": name, "role": role, "user": user, "allow": allow}


class TestAccessRules(IntegrationTestCase):
    def allowed(self, rules, roles=("Sales User",), doctype=access.SKILL):
        with patch.object(access, "_rules", return_value=rules), patch.object(access.frappe, "get_roles", return_value=list(roles)):
            return access.allowed_names(doctype, USER)

    def test_nothing_is_allowed_without_a_rule(self):
        self.assertEqual(self.allowed([]), set())

    def test_everyone_row_allows(self):
        self.assertEqual(self.allowed([rule("sales")]), {"sales"})

    def test_role_row_needs_the_role(self):
        self.assertEqual(self.allowed([rule("sales", role="Sales User")]), {"sales"})
        self.assertEqual(self.allowed([rule("sales", role="Accounts User")]), set())

    def test_user_row_beats_role_row(self):
        rules = [rule("sales", role="Sales User", allow=0), rule("sales", user=USER)]
        self.assertEqual(self.allowed(rules), {"sales"})
        rules = [rule("sales", role="Sales User"), rule("sales", user=USER, allow=0)]
        self.assertEqual(self.allowed(rules), set())

    def test_role_row_beats_everyone_row(self):
        self.assertEqual(self.allowed([rule("sales", allow=0), rule("sales", role="Sales User")]), {"sales"})

    def test_deny_wins_on_the_same_level(self):
        rules = [rule("sales", role="Sales User"), rule("sales", role="Stock User", allow=0)]
        self.assertEqual(self.allowed(rules, roles=("Sales User", "Stock User")), set())

    def test_skills_and_plugins_are_separate(self):
        rules = [rule("excel", doctype=access.PLUGIN)]
        self.assertEqual(self.allowed(rules, doctype=access.SKILL), set())
        self.assertEqual(self.allowed(rules, doctype=access.PLUGIN), {"excel"})

    def test_administrator_is_always_allowed(self):
        with patch.object(access, "_rules", return_value=[]):
            self.assertIsNone(access.allowed_names(access.SKILL, "Administrator"))

    def test_plugin_tool_follows_its_plugin(self):
        with patch.object(access, "plugin_tools", return_value={"export.excel": ["excel"]}), \
                patch.object(access, "_rules", return_value=[]), patch.object(access.frappe, "get_roles", return_value=[]):
            self.assertFalse(access.tool_allowed("export.excel", USER))
            self.assertTrue(access.tool_allowed("data.query", USER))   # not a plugin tool
        with patch.object(access, "plugin_tools", return_value={"export.excel": ["excel"]}), \
                patch.object(access, "_rules", return_value=[rule("excel", doctype=access.PLUGIN)]), patch.object(access.frappe, "get_roles", return_value=[]):
            self.assertTrue(access.tool_allowed("export.excel", USER))

    def test_settings_reject_a_row_with_both_role_and_user(self):
        settings = frappe.get_single("AI Settings")
        settings.append("access_rules", {"reference_doctype": "AI Plugin", "reference_name": "excel", "role": "System Manager", "user": "Administrator", "allow": 1})
        self.assertRaises(frappe.ValidationError, settings.validate)


class TestExportTools(IntegrationTestCase):
    ROWS = [
        {"customer": "عميل أ", "net_total": 1500.5, "qty": 3},
        {"customer": "Customer B", "net_total": 820, "qty": 1},
    ]

    def setUp(self):
        frappe.set_user("Administrator")
        self.run = make_test_run()
        self.ctx = frappe._dict(user="Administrator", run_name=self.run.name)

    def tearDown(self):
        delete_docs("File", frappe.get_all("File", filters={"attached_to_doctype": "AI Run", "attached_to_name": self.run.name}, pluck="name"))
        delete_docs("AI Run", [self.run.name])

    def assert_file(self, result, ext):
        self.assertNotIn("error", result, result)
        self.assertTrue(result["file_name"].endswith("." + ext))
        doc = frappe.get_doc("File", {"file_url": result["file_url"]})
        self.assertTrue(doc.is_private)
        self.assertEqual((doc.attached_to_doctype, doc.attached_to_name), ("AI Run", self.run.name))
        self.assertTrue(os.path.getsize(doc.get_full_path()) > 100)
        return doc

    def test_excel(self):
        from openpyxl import load_workbook

        from ai_agent.tools import exports

        doc = self.assert_file(exports.handle_excel({"title": "مبيعات العملاء", "rows": self.ROWS}, self.ctx), "xlsx")
        ws = load_workbook(doc.get_full_path()).active
        self.assertTrue(ws.sheet_view.rightToLeft)
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(2, 2).value, 1500.5)   # numbers stay numbers

    def test_csv_has_bom_for_arabic(self):
        from ai_agent.tools import exports

        doc = self.assert_file(exports.handle_csv({"title": "x", "rows": self.ROWS}, self.ctx), "csv")
        with open(doc.get_full_path(), "rb") as f:
            self.assertTrue(f.read().startswith(b"\xef\xbb\xbf"))

    def test_pdf(self):
        from ai_agent.tools import exports

        doc = self.assert_file(exports.handle_pdf({"title": "تقرير", "summary": "سطر أول\nسطر ثانٍ", "rows": self.ROWS}, self.ctx), "pdf")
        with open(doc.get_full_path(), "rb") as f:
            self.assertEqual(f.read(4), b"%PDF")

    def test_image_chart_and_table(self):
        from ai_agent.tools import exports

        chart = {"type": "bar", "labels": ["يناير", "فبراير"], "datasets": [{"name": "المبيعات", "values": [10, 20]}]}
        result = exports.handle_image({"title": "المبيعات الشهرية", "chart": chart}, self.ctx)
        self.assertTrue(result.get("image"))
        self.assert_file(result, "png")
        self.assert_file(exports.handle_image({"title": "جدول", "rows": self.ROWS}, self.ctx), "png")

    def test_artifact_of_another_user_is_refused(self):
        from ai_agent.runtime.artifact_service import ArtifactService
        from ai_agent.tools import exports

        other = make_test_run(user="Guest")
        try:
            artifact = ArtifactService().create_table(other.name, "t", self.ROWS)
            result = exports.handle_excel({"artifact": artifact.name}, self.ctx)
            self.assertIn("error", result)
        finally:
            delete_docs("AI Run", [other.name])

    def test_empty_source_is_an_error(self):
        from ai_agent.tools import exports

        self.assertIn("error", exports.handle_excel({"rows": []}, self.ctx))


class TestPrintAndStatement(IntegrationTestCase):
    def setUp(self):
        frappe.set_user("Administrator")
        self.run = make_test_run()
        self.ctx = frappe._dict(user="Administrator", run_name=self.run.name)

    def tearDown(self):
        delete_docs("File", frappe.get_all("File", filters={"attached_to_doctype": "AI Run", "attached_to_name": self.run.name}, pluck="name"))
        delete_docs("AI Run", [self.run.name])

    def test_print_document_as_pdf(self):
        from ai_agent.tools import exports

        name = frappe.db.get_value("Sales Invoice", {"docstatus": 1}, "name")
        if not name:
            self.skipTest("no submitted Sales Invoice on this site")
        result = exports.handle_print({"doctype": "Sales Invoice", "name": name}, self.ctx)
        self.assertNotIn("error", result, result)
        path = frappe.get_doc("File", {"file_url": result["file_url"]}).get_full_path()
        with open(path, "rb") as f:
            self.assertEqual(f.read(4), b"%PDF")

    def test_print_missing_document(self):
        from ai_agent.tools import exports

        self.assertIn("error", exports.handle_print({"doctype": "Sales Invoice", "name": "no-such-invoice"}, self.ctx))

    def test_statement_balances_match_the_ledger(self):
        from ai_agent.tools import exports

        row = frappe.db.sql("select party, company from `tabGL Entry` where party_type='Customer' and is_cancelled=0 limit 1", as_dict=True)
        if not row:
            self.skipTest("no customer GL entries on this site")
        party, company = row[0].party, row[0].company
        result = exports.handle_statement({"party_type": "Customer", "party": party, "company": company, "from_date": "2000-01-01", "to_date": "2099-12-31"}, self.ctx)
        self.assertNotIn("error", result, result)
        ledger = frappe.db.sql("select coalesce(sum(debit - credit), 0) from `tabGL Entry` where party_type='Customer' and party=%s and company=%s and is_cancelled=0", (party, company))[0][0]
        self.assertAlmostEqual(float(result["closing"]), float(ledger), places=2)


class TestEmailPlugin(IntegrationTestCase):
    def setUp(self):
        frappe.set_user("Administrator")
        self.run = make_test_run()
        self.ctx = frappe._dict(user="Administrator", run_name=self.run.name)

    def tearDown(self):
        delete_docs("AI Approval Request", frappe.get_all("AI Approval Request", filters={"run": self.run.name}, pluck="name"))
        delete_docs("AI Run", [self.run.name])

    def payload(self, **extra):
        from ai_agent.tools import messaging

        return messaging.build_payload({"recipients": ["a@example.com", "bad-address"], "subject": "كشف حساب", "message": "مرحباً\nمرفق الكشف", **extra})

    def test_payload_keeps_valid_addresses_only(self):
        self.assertEqual(self.payload()["recipients"], ["a@example.com"])
        from ai_agent.tools import messaging

        self.assertRaises(frappe.ValidationError, messaging.build_payload, {"recipients": ["nope"], "subject": "s", "message": "m"})

    def test_request_needs_an_outgoing_account(self):
        from ai_agent.runtime.approval_engine import ApprovalEngine
        from ai_agent.runtime.exceptions import PermissionDenied
        from ai_agent.tools import messaging

        with patch.object(messaging, "outgoing_account_ready", return_value=False):
            self.assertRaises(PermissionDenied, ApprovalEngine(self.ctx).create_approval_request,
                              run_name=self.run.name, action_type="Approved Method", action_summary="x", payload=self.payload())

    def test_unknown_method_is_refused(self):
        from ai_agent.runtime.approval_engine import ApprovalEngine
        from ai_agent.runtime.exceptions import PermissionDenied

        self.assertRaises(PermissionDenied, ApprovalEngine(self.ctx).create_approval_request,
                          run_name=self.run.name, action_type="Approved Method", action_summary="x", payload={"method": "frappe.delete_doc"})

    def test_sent_only_after_approval(self):
        from ai_agent.runtime.approval_engine import ApprovalEngine
        from ai_agent.tools import messaging

        sent = []
        with patch.object(messaging, "outgoing_account_ready", return_value=True), \
                patch("frappe.core.doctype.communication.email.make", side_effect=lambda **kw: sent.append(kw) or {"name": "COMM-TEST"}):
            engine = ApprovalEngine(self.ctx)
            approval = engine.create_approval_request(run_name=self.run.name, action_type="Approved Method",
                                                      action_summary=messaging.summary(self.payload()), payload=self.payload())
            self.assertEqual(sent, [])                       # requesting sends nothing
            with patch.object(frappe.get_cached_doc("AI Settings"), "require_four_eyes", 0):
                engine.approve(approval.name)
                self.assertEqual(sent, [])                   # approving records the decision
                result = engine.execute_approved_action(approval.name)
        self.assertEqual(len(sent), 1)
        self.assertEqual(sent[0]["recipients"], "a@example.com")
        self.assertIn("<br>", sent[0]["content"])
        self.assertEqual(result["communication"], "COMM-TEST")

    def test_another_users_file_cannot_be_attached(self):
        from ai_agent.tools import messaging

        f = frappe.get_doc({"doctype": "File", "file_name": "secret.txt", "content": b"x", "is_private": 1}).insert(ignore_permissions=True)
        frappe.db.set_value("File", f.name, "owner", "Guest")
        try:
            with patch.object(messaging, "outgoing_account_ready", return_value=True), \
                    patch.object(messaging.frappe, "has_permission", side_effect=lambda dt, *a, **k: dt != "File"):
                self.assertRaises(frappe.PermissionError, messaging.assert_email_allowed, self.payload(attachments=[f.file_url]), "Administrator")
        finally:
            delete_docs("File", [f.name])
